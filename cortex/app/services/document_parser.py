"""
Document Parser Service
Extracts text from various document formats
"""
import os
import io
from typing import List, Optional, Dict, Any, Tuple
from dataclasses import dataclass
from enum import Enum
import mimetypes

from app.core.config import settings


class DocumentFormat(str, Enum):
    """Supported document formats."""
    PDF = "pdf"
    DOCX = "docx"
    TXT = "txt"
    MD = "md"
    HTML = "html"
    CSV = "csv"
    JSON = "json"


@dataclass
class ParsedDocument:
    """Result of document parsing."""
    content: str
    format: DocumentFormat
    metadata: Dict[str, Any]
    page_count: int
    char_count: int
    word_count: int
    
    def to_dict(self) -> dict:
        return {
            "content": self.content,
            "format": self.format.value,
            "metadata": self.metadata,
            "page_count": self.page_count,
            "char_count": self.char_count,
            "word_count": self.word_count,
        }


class DocumentParser:
    """
    Parses various document formats to extract text.
    
    Supports:
    - PDF (via pypdf)
    - DOCX (via python-docx)
    - TXT, MD (direct read)
    - HTML (via BeautifulSoup)
    - CSV, JSON (structured extraction)
    """
    
    MIME_TO_FORMAT = {
        "application/pdf": DocumentFormat.PDF,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": DocumentFormat.DOCX,
        "text/plain": DocumentFormat.TXT,
        "text/markdown": DocumentFormat.MD,
        "text/html": DocumentFormat.HTML,
        "text/csv": DocumentFormat.CSV,
        "application/json": DocumentFormat.JSON,
    }
    
    EXT_TO_FORMAT = {
        ".pdf": DocumentFormat.PDF,
        ".docx": DocumentFormat.DOCX,
        ".txt": DocumentFormat.TXT,
        ".md": DocumentFormat.MD,
        ".markdown": DocumentFormat.MD,
        ".html": DocumentFormat.HTML,
        ".htm": DocumentFormat.HTML,
        ".csv": DocumentFormat.CSV,
        ".json": DocumentFormat.JSON,
    }
    
    def detect_format(
        self,
        filename: str = None,
        mime_type: str = None,
    ) -> Optional[DocumentFormat]:
        """Detect document format from filename or MIME type."""
        if mime_type and mime_type in self.MIME_TO_FORMAT:
            return self.MIME_TO_FORMAT[mime_type]
        
        if filename:
            ext = os.path.splitext(filename)[1].lower()
            if ext in self.EXT_TO_FORMAT:
                return self.EXT_TO_FORMAT[ext]
        
        return None
    
    def parse(
        self,
        content: bytes,
        filename: str = None,
        mime_type: str = None,
    ) -> ParsedDocument:
        """
        Parse document content and extract text.
        
        Args:
            content: Raw file bytes
            filename: Original filename (for format detection)
            mime_type: MIME type (for format detection)
            
        Returns:
            ParsedDocument with extracted text and metadata
        """
        format = self.detect_format(filename, mime_type)
        
        if format is None:
            # Try to detect from content
            format = self._detect_from_content(content)
        
        if format is None:
            raise ValueError(f"Unsupported document format: {filename or mime_type}")
        
        # Parse based on format
        if format == DocumentFormat.PDF:
            return self._parse_pdf(content, filename)
        elif format == DocumentFormat.DOCX:
            return self._parse_docx(content, filename)
        elif format in [DocumentFormat.TXT, DocumentFormat.MD]:
            return self._parse_text(content, filename, format)
        elif format == DocumentFormat.HTML:
            return self._parse_html(content, filename)
        elif format == DocumentFormat.CSV:
            return self._parse_csv(content, filename)
        elif format == DocumentFormat.JSON:
            return self._parse_json(content, filename)
        else:
            raise ValueError(f"Parser not implemented for format: {format}")
    
    def _detect_from_content(self, content: bytes) -> Optional[DocumentFormat]:
        """Try to detect format from file signature."""
        # PDF magic bytes
        if content[:4] == b'%PDF':
            return DocumentFormat.PDF
        
        # DOCX (ZIP with specific structure)
        if content[:4] == b'PK\x03\x04':
            return DocumentFormat.DOCX
        
        # Try as text
        try:
            text = content.decode('utf-8')
            if text.strip().startswith('<'):
                return DocumentFormat.HTML
            if text.strip().startswith('{') or text.strip().startswith('['):
                return DocumentFormat.JSON
            return DocumentFormat.TXT
        except UnicodeDecodeError:
            pass
        
        return None
    
    def _parse_pdf(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse PDF document."""
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ImportError("pypdf package not installed")
        
        reader = PdfReader(io.BytesIO(content))
        
        text_parts = []
        for page in reader.pages:
            text = page.extract_text() or ""
            text_parts.append(text)
        
        full_text = "\n\n".join(text_parts)
        
        # Extract metadata
        metadata = {}
        if reader.metadata:
            metadata = {
                "title": reader.metadata.title,
                "author": reader.metadata.author,
                "subject": reader.metadata.subject,
                "creator": reader.metadata.creator,
            }
        
        return ParsedDocument(
            content=full_text,
            format=DocumentFormat.PDF,
            metadata=metadata,
            page_count=len(reader.pages),
            char_count=len(full_text),
            word_count=len(full_text.split()),
        )
    
    def _parse_docx(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse DOCX document."""
        try:
            from docx import Document
        except ImportError:
            raise ImportError("python-docx package not installed")
        
        doc = Document(io.BytesIO(content))
        
        text_parts = []
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells)
                if row_text.strip():
                    text_parts.append(row_text)
        
        full_text = "\n\n".join(text_parts)
        
        # Extract metadata
        metadata = {}
        if doc.core_properties:
            metadata = {
                "title": doc.core_properties.title,
                "author": doc.core_properties.author,
                "subject": doc.core_properties.subject,
                "created": str(doc.core_properties.created) if doc.core_properties.created else None,
            }
        
        return ParsedDocument(
            content=full_text,
            format=DocumentFormat.DOCX,
            metadata=metadata,
            page_count=1,  # DOCX doesn't have pages in the same way
            char_count=len(full_text),
            word_count=len(full_text.split()),
        )
    
    def _parse_text(
        self,
        content: bytes,
        filename: str,
        format: DocumentFormat,
    ) -> ParsedDocument:
        """Parse plain text or markdown."""
        # Try different encodings
        for encoding in ['utf-8', 'utf-16', 'latin-1']:
            try:
                text = content.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            text = content.decode('utf-8', errors='replace')
        
        return ParsedDocument(
            content=text,
            format=format,
            metadata={"filename": filename},
            page_count=1,
            char_count=len(text),
            word_count=len(text.split()),
        )
    
    def _parse_html(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse HTML document."""
        try:
            from bs4 import BeautifulSoup
        except ImportError:
            # Fallback: strip tags with regex
            import re
            text = content.decode('utf-8', errors='replace')
            text = re.sub(r'<script[^>]*>.*?</script>', '', text, flags=re.DOTALL)
            text = re.sub(r'<style[^>]*>.*?</style>', '', text, flags=re.DOTALL)
            text = re.sub(r'<[^>]+>', ' ', text)
            text = re.sub(r'\s+', ' ', text).strip()
            
            return ParsedDocument(
                content=text,
                format=DocumentFormat.HTML,
                metadata={"filename": filename},
                page_count=1,
                char_count=len(text),
                word_count=len(text.split()),
            )
        
        soup = BeautifulSoup(content, 'html.parser')
        
        # Remove script and style elements
        for element in soup(['script', 'style', 'nav', 'footer', 'header']):
            element.decompose()
        
        text = soup.get_text(separator='\n', strip=True)
        
        # Extract metadata
        metadata = {"filename": filename}
        title_tag = soup.find('title')
        if title_tag:
            metadata["title"] = title_tag.get_text()
        
        return ParsedDocument(
            content=text,
            format=DocumentFormat.HTML,
            metadata=metadata,
            page_count=1,
            char_count=len(text),
            word_count=len(text.split()),
        )
    
    def _parse_csv(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse CSV document."""
        import csv
        
        text = content.decode('utf-8', errors='replace')
        reader = csv.reader(io.StringIO(text))
        
        rows = []
        for row in reader:
            rows.append(" | ".join(row))
        
        full_text = "\n".join(rows)
        
        return ParsedDocument(
            content=full_text,
            format=DocumentFormat.CSV,
            metadata={"filename": filename, "row_count": len(rows)},
            page_count=1,
            char_count=len(full_text),
            word_count=len(full_text.split()),
        )
    
    def _parse_json(self, content: bytes, filename: str) -> ParsedDocument:
        """Parse JSON document."""
        import json
        
        text = content.decode('utf-8', errors='replace')
        
        try:
            data = json.loads(text)
            # Pretty print for readability
            full_text = json.dumps(data, indent=2, ensure_ascii=False)
        except json.JSONDecodeError:
            full_text = text
        
        return ParsedDocument(
            content=full_text,
            format=DocumentFormat.JSON,
            metadata={"filename": filename},
            page_count=1,
            char_count=len(full_text),
            word_count=len(full_text.split()),
        )


# Singleton instance
_parser: Optional[DocumentParser] = None


def get_document_parser() -> DocumentParser:
    """Get or create the document parser instance."""
    global _parser
    if _parser is None:
        _parser = DocumentParser()
    return _parser
